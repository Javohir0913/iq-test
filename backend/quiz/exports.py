"""Urinishlarni CSV va Excel (.xlsx) ga eksport qilish."""

import csv
import io

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import CATEGORY_NAMES

HEADERS = [
    "ID",
    "Ism familiya",
    "Telefon",
    "Telegram ID",
    "Til",
    "Bosqich",
    "Holat",
    "Boshlangan",
    "Tugagan",
    "Davomiylik (daq)",
    "Javoblar",
    "Ball",
    "Foiz",
    "Daraja",
] + [f"{names['uz']} (%)" for names in CATEGORY_NAMES.values()]


def _fmt(dt):
    return timezone.localtime(dt).strftime("%Y-%m-%d %H:%M:%S") if dt else ""


def attempt_rows(queryset):
    qs = queryset.select_related("participant", "level", "stage").prefetch_related("answers")
    for a in qs:
        p = a.participant
        duration = round((a.finished_at - a.started_at).total_seconds() / 60, 1) if a.finished_at else ""
        cats = a.category_scores or {}
        yield [
            str(a.id),
            p.full_name,
            p.phone,
            p.telegram_id or "",
            p.language,
            a.stage.number if a.stage else "",
            a.get_status_display(),
            _fmt(a.started_at),
            _fmt(a.finished_at),
            duration,
            len(a.answers.all()),
            a.total_score if a.total_score is not None else "",
            float(a.percent) if a.percent is not None else "",
            a.level.name_uz if a.level else "",
        ] + [cats.get(key, {}).get("percent", "") for key in CATEGORY_NAMES]


def export_csv(queryset):
    buf = io.StringIO()
    buf.write("﻿")  # Excel uchun UTF-8 BOM
    writer = csv.writer(buf)
    writer.writerow(HEADERS)
    for row in attempt_rows(queryset):
        writer.writerow(row)
    response = HttpResponse(buf.getvalue(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="urinishlar-{timezone.localdate():%Y%m%d}.csv"'
    return response


def export_xlsx(queryset):
    wb = Workbook()
    ws = wb.active
    ws.title = "Urinishlar"
    ws.append(HEADERS)
    header_fill = PatternFill("solid", fgColor="1B2A41")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    for row in attempt_rows(queryset):
        ws.append(row)
    for idx, header in enumerate(HEADERS, start=1):
        width = max(len(str(header)), *(len(str(c.value or "")) for c in ws[get_column_letter(idx)][1:200]), 8)
        ws.column_dimensions[get_column_letter(idx)].width = min(width + 2, 42)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(
        buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="urinishlar-{timezone.localdate():%Y%m%d}.xlsx"'
    return response
